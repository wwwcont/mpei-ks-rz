"""Оформление docx: Times New Roman 14, 1,5 интервала, таблицы и рисунки с подписями, содержание."""
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Times New Roman"


def _font(style_or_run, size=14, bold=None):
    f = style_or_run.font
    f.name = FONT
    f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    f.color.rgb = RGBColor(0, 0, 0)
    el = style_or_run.element if hasattr(style_or_run, "element") else style_or_run._element
    rpr = el.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rfonts.get(qn(a)) is not None:
            del rfonts.attrib[qn(a)]
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(a), FONT)


class Doc:
    def __init__(self):
        self.d = Document()
        self.tables = 0
        self.figures = 0
        s = self.d.sections[0]
        s.page_height, s.page_width = Cm(29.7), Cm(21.0)
        s.left_margin, s.right_margin = Cm(3), Cm(1.5)
        s.top_margin, s.bottom_margin = Cm(2), Cm(2)
        normal = self.d.styles["Normal"]
        _font(normal)
        pf = normal.paragraph_format
        pf.line_spacing = 1.5
        pf.space_after = Pt(0)
        pf.space_before = Pt(0)
        pf.first_line_indent = Cm(1.25)
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for name, size in (("Heading 1", 14), ("Heading 2", 14)):
            st = self.d.styles[name]
            _font(st, size, bold=True)
            st.font.italic = False
            st.paragraph_format.space_before = Pt(0)
            st.paragraph_format.space_after = Pt(12)
            st.paragraph_format.keep_with_next = True
            st.paragraph_format.line_spacing = 1.5
        h1 = self.d.styles["Heading 1"].paragraph_format
        h1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        h1.first_line_indent = Cm(0)
        self.d.styles["Heading 1"].font.all_caps = False
        h2 = self.d.styles["Heading 2"].paragraph_format
        h2.first_line_indent = Cm(1.25)
        h2.space_before = Pt(12)

    # ---- текст ----
    def p(self, text="", bold=False, align=None, indent=True, size=14, keep=False, italic=False):
        par = self.d.add_paragraph()
        if align == "center":
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "right":
            par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif align == "left":
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if not indent:
            par.paragraph_format.first_line_indent = Cm(0)
        if keep:
            par.paragraph_format.keep_with_next = True
        self._runs(par, text, bold, size, italic)
        return par

    def _runs(self, par, text, bold=False, size=14, italic=False):
        # **жирный** внутри строки
        parts = text.split("**")
        for i, t in enumerate(parts):
            if not t:
                continue
            r = par.add_run(t)
            _font(r, size, bold=bold or i % 2 == 1)
            r.italic = italic

    def bullets(self, items):
        for it in items:
            par = self.p("", indent=True)
            self._runs(par, "– " + it)

    def h1(self, text, new_page=True):
        if new_page:
            self.page_break()
        return self.d.add_heading(text, level=1)

    def h2(self, text):
        return self.d.add_heading(text, level=2)

    def page_break(self):
        par = self.d.add_paragraph()
        par.paragraph_format.first_line_indent = Cm(0)
        par.add_run().add_break(WD_BREAK.PAGE)

    # ---- таблицы ----
    def table(self, title, headers, rows, widths=None, size=12, align=None, head_size=None):
        """rows: кортежи значений или ("section", "текст") — объединённая строка-подзаголовок."""
        self.tables += 1
        cap = self.p(f"Таблица {self.tables}", indent=False, align="right", keep=True)
        cap.paragraph_format.space_before = Pt(6)
        cap.paragraph_format.line_spacing = 1.0
        name = self.p(title, indent=False, align="center", keep=True)
        name.paragraph_format.line_spacing = 1.0
        name.paragraph_format.space_after = Pt(3)
        ncol = len(headers)
        t = self.d.add_table(rows=1, cols=ncol)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(headers):
            self._cell(t.rows[0].cells[i], h, head_size or size, bold=True)
        _repeat_header(t.rows[0])
        for r in rows:
            row = t.add_row()
            _cant_split(row)
            if r and r[0] == "section":
                c = row.cells[0].merge(row.cells[-1])
                self._cell(c, r[1], size, italic=True)
                continue
            for i, val in enumerate(r):
                a = align[i] if align else "center"
                self._cell(row.cells[i], str(val), size, align=a)
        widths = widths or [16.5 / ncol] * ncol
        k = 16.5 / sum(widths) if sum(widths) > 16.5 else 1
        widths = [w * k for w in widths]
        t.autofit = False
        tblpr = t._tbl.tblPr
        lay = OxmlElement("w:tblLayout")
        lay.set(qn("w:type"), "fixed")
        tblpr.append(lay)
        for i, w in enumerate(widths):
            t.columns[i].width = Cm(w)
        for row in t.rows:
            for i, w in enumerate(widths):
                if i < len(row.cells):
                    row.cells[i].width = Cm(w)
        # короткую таблицу не разрывать между страницами; у длинной — не оставлять «хвост» из 1–3 строк
        for row in (t.rows[:-1] if len(rows) <= 20 else t.rows[-5:-1]):
            for c in row.cells:
                for par in c.paragraphs:
                    par.paragraph_format.keep_with_next = True
        self.p("", indent=False).paragraph_format.line_spacing = 1.0
        return t

    def _cell(self, cell, text, size, bold=False, italic=False, align="center"):
        cell.text = ""
        par = cell.paragraphs[0]
        par.paragraph_format.first_line_indent = Cm(0)
        par.paragraph_format.line_spacing = 1.0
        par.alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER, "left": WD_ALIGN_PARAGRAPH.LEFT,
                         "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
        r = par.add_run(text)
        _font(r, size, bold=bold)
        r.italic = italic

    # ---- рисунки ----
    def figure(self, path, title, width_cm=16.0, max_h_cm=21.0):
        from PIL import Image
        w, h = Image.open(path).size
        width_cm = min(width_cm, max_h_cm * w / h)
        self.figures += 1
        par = self.d.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.first_line_indent = Cm(0)
        par.paragraph_format.keep_with_next = True
        par.paragraph_format.line_spacing = 1.0
        par.paragraph_format.space_before = Pt(6)
        par.add_run().add_picture(str(path), width=Cm(width_cm))
        self.p(f"Рис. {self.figures}. {title}", align="center", indent=False)
        return self.figures

    def landscape(self, on=True):
        """Новый раздел книжной/альбомной ориентации (колонтитулы — как в предыдущем)."""
        from docx.enum.section import WD_ORIENT
        sec = self.d.add_section(WD_SECTION.NEW_PAGE)
        sec.different_first_page_header_footer = False
        if on:
            sec.orientation = WD_ORIENT.LANDSCAPE
            sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
            sec.left_margin, sec.right_margin = Cm(2), Cm(2)
            sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5)
        else:
            sec.orientation = WD_ORIENT.PORTRAIT
            sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
            sec.left_margin, sec.right_margin = Cm(3), Cm(1.5)
            sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)
        return sec

    def kv_table(self, rows, indent_cm=3.6, col1=4.1, col2=8.8):
        """Таблица без рамок для титула: [(подпись, значение)], подпись жирная; ("Выполнил", None) — заголовок."""
        t = self.d.add_table(rows=0, cols=2)
        t.autofit = False
        tblpr = t._tbl.tblPr
        ind = OxmlElement("w:tblInd")
        ind.set(qn("w:w"), str(int(indent_cm * 567)))
        ind.set(qn("w:type"), "dxa")
        tblpr.append(ind)
        lay = OxmlElement("w:tblLayout")
        lay.set(qn("w:type"), "fixed")
        tblpr.append(lay)
        t.columns[0].width, t.columns[1].width = Cm(col1), Cm(col2)
        for k, v in rows:
            r = t.add_row()
            r.cells[0].width, r.cells[1].width = Cm(col1), Cm(col2)
            if v is None:
                c = r.cells[0].merge(r.cells[1])
                self._cell(c, k, 14, bold=True, align="left")
            else:
                self._cell(r.cells[0], k, 14, bold=True, align="left")
                self._cell(r.cells[1], v, 14, align="left")
        return t

    def figure_best(self, path, title, page_w_cm=16.5, page_h_cm=23.0):
        """Рисунок крупнее всего: как есть по ширине страницы или повёрнутым на всю страницу."""
        from PIL import Image
        w, h = Image.open(path).size
        upright = min(page_w_cm / w, (page_h_cm - 1.5) / h)
        rotated = min(page_h_cm / w, (page_w_cm - 1.2) / h)
        if rotated > upright * 1.1:
            return self.figure_rotated(path, title, page_w_cm, page_h_cm)
        return self.figure(path, title, page_w_cm, page_h_cm - 1.5)

    @staticmethod
    def best_scale(path, page_w_cm=16.5, page_h_cm=23.0):
        from PIL import Image
        w, h = Image.open(path).size
        return max(min(page_w_cm / w, (page_h_cm - 1.5) / h), min(page_h_cm / w, (page_w_cm - 1.2) / h))

    def figure_rotated(self, path, title, page_w_cm=16.5, page_h_cm=23.5):
        """Широкий рисунок на всю книжную страницу, повёрнутый на 90° (читается как альбомный);
        подпись «Рис. N. …» впечатана в изображение и повёрнута вместе с ним."""
        from PIL import Image, ImageDraw, ImageFont
        from .common import ROOT
        self.figures += 1
        im = Image.open(path).convert("RGB")
        w, h = im.size
        # после поворота ширина рисунка ляжет вдоль высоты страницы
        scale_cm = min(page_h_cm / w, (page_w_cm - 1.2) / h)       # см на пиксель
        px14 = int(0.494 / scale_cm)                              # 14 pt в пикселях рисунка
        f = ImageFont.truetype(str(ROOT / "fonts" / "LiberationSerif-Regular.ttf"), px14)
        cap = f"Рис. {self.figures}. {title}"
        strip = int(px14 * 2.2)
        out = Image.new("RGB", (w, h + strip), "white")
        out.paste(im, (0, 0))
        ImageDraw.Draw(out).text((w / 2, h + strip / 2), cap, font=f, fill=(0, 0, 0), anchor="mm")
        out = out.rotate(90, expand=True)
        rot = Path(path).with_name(Path(path).stem + "_rot.png")
        out.save(rot)
        par = self.d.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.first_line_indent = Cm(0)
        par.paragraph_format.line_spacing = 1.0
        wr, hr = out.size
        width = min(page_w_cm, page_h_cm * wr / hr)
        par.add_run().add_picture(str(rot), width=Cm(width))
        return self.figures

    # ---- служебное ----
    def toc(self):
        self.p("СОДЕРЖАНИЕ", bold=True, align="center", indent=False)
        par = self.d.add_paragraph()
        par.paragraph_format.first_line_indent = Cm(0)
        _field(par, 'TOC \\o "1-2" \\h \\z \\u', "Содержание обновится при открытии файла (F9 — обновить поле).")

    def page_numbers(self, first_footer=None, first_header=()):
        s = self.d.sections[0]
        s.different_first_page_header_footer = True
        hp = s.first_page_header.paragraphs[0]
        for i, line in enumerate(first_header):
            par = hp if i == 0 else s.first_page_header.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.first_line_indent = Cm(0)
            par.paragraph_format.line_spacing = 1.0
            self._runs(par, line)
        if first_footer:
            fp = s.first_page_footer.paragraphs[0]
            fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fp.paragraph_format.first_line_indent = Cm(0)
            self._runs(fp, first_footer)
        par = s.footer.paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.first_line_indent = Cm(0)
        _field(par, "PAGE", "1")

    def save(self, path):
        self.d.save(str(path))


def _field(par, instr, placeholder):
    def run_with(el):
        r = par.add_run()
        _font(r, 14)
        r._r.append(el)
        return r
    b = OxmlElement("w:fldChar")
    b.set(qn("w:fldCharType"), "begin")
    run_with(b)
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    run_with(it)
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run_with(sep)
    r = par.add_run(placeholder)
    _font(r, 14)
    e = OxmlElement("w:fldChar")
    e.set(qn("w:fldCharType"), "end")
    run_with(e)


def _repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def _cant_split(row):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    trpr.append(el)
